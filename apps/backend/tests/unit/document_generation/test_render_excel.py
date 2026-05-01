"""
Tests for the Excel (.xlsx) renderer.
Uses in-memory bytes buffers — no disk writes.
"""
import io

import openpyxl

from docuweave.use_cases.document_generation.render_excel import render_excel
from docuweave.domain.entities.template import Template, TemplateComponent


def _parse_xlsx(buffer: io.BytesIO) -> openpyxl.Workbook:
    buffer.seek(0)
    return openpyxl.load_workbook(buffer)


def test_headers_in_first_row() -> None:
    template = Template(
        id="t1",
        name="Test",
        output_format="xlsx",
        components=(
            TemplateComponent(
                id="c1",
                type="table",
                position={"x": 0, "y": 0, "width": 500, "height": "auto"},
                source_table="users",
                columns=(
                    {"header": "Name", "field": "name"},
                    {"header": "Email", "field": "email"},
                ),
            ),
        ),
    )
    data = {"users": [{"name": "Alice", "email": "alice@test.com"}]}
    buf = render_excel(template, data=data)
    wb = _parse_xlsx(buf)
    ws = wb.active
    assert ws is not None
    assert ws.cell(1, 1).value == "Name"
    assert ws.cell(1, 2).value == "Email"


def test_data_rows_populated() -> None:
    template = Template(
        id="t2",
        name="Test",
        output_format="xlsx",
        components=(
            TemplateComponent(
                id="c2",
                type="table",
                position={"x": 0, "y": 0, "width": 500, "height": "auto"},
                source_table="users",
                columns=({"header": "Name", "field": "name"},),
            ),
        ),
    )
    data = {"users": [{"name": "Alice"}, {"name": "Bob"}]}
    buf = render_excel(template, data=data)
    wb = _parse_xlsx(buf)
    ws = wb.active
    assert ws is not None
    names = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    assert "Alice" in names
    assert "Bob" in names


def test_text_component_written_to_cell() -> None:
    template = Template(
        id="t3",
        name="Test",
        output_format="xlsx",
        components=(
            TemplateComponent(
                id="c3",
                type="text",
                position={"x": 0, "y": 0, "width": 300, "height": "auto"},
                content="Department Report",
            ),
        ),
    )
    buf = render_excel(template, data={})
    wb = _parse_xlsx(buf)
    ws = wb.active
    assert ws is not None
    all_values = [ws.cell(r, c).value for r in range(1, ws.max_row + 1) for c in range(1, 10)]
    assert "Department Report" in all_values
