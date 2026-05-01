"""
Tests for the Word (.docx) renderer.
Uses in-memory bytes buffers — no disk writes.
"""
import io

from docx import Document as DocxDocument

from docuweave.use_cases.document_generation.render_word import render_word
from docuweave.domain.entities.template import Template, TemplateComponent


def _parse_docx(buffer: io.BytesIO) -> DocxDocument:
    buffer.seek(0)
    return DocxDocument(buffer)


def _all_text(doc: DocxDocument) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def test_text_component_appears_in_docx() -> None:
    template = Template(
        id="t1",
        name="Test",
        output_format="docx",
        components=(
            TemplateComponent(
                id="c1",
                type="text",
                position={"x": 0, "y": 0, "width": 300, "height": "auto"},
                content="Hello DocuWeave",
            ),
        ),
    )
    buf = render_word(template, data={})
    doc = _parse_docx(buf)
    assert "Hello DocuWeave" in _all_text(doc)


def test_dynamic_field_substitutes_value() -> None:
    template = Template(
        id="t2",
        name="Test",
        output_format="docx",
        components=(
            TemplateComponent(
                id="c2",
                type="dynamic_field",
                position={"x": 0, "y": 0, "width": 300, "height": "auto"},
                source_table="users",
                source_field="name",
            ),
        ),
    )
    data = {"users": [{"name": "Alice"}]}
    buf = render_word(template, data=data)
    doc = _parse_docx(buf)
    assert "Alice" in _all_text(doc)


def test_table_component_renders_headers() -> None:
    template = Template(
        id="t3",
        name="Test",
        output_format="docx",
        components=(
            TemplateComponent(
                id="c3",
                type="table",
                position={"x": 0, "y": 0, "width": 500, "height": "auto"},
                source_table="users",
                columns=(
                    {"header": "Name", "field": "name"},
                    {"header": "Email", "field": "email"},
                ),
            ),
        ),
    )
    data = {"users": [{"name": "Alice", "email": "alice@test.com"}]}
    buf = render_word(template, data=data)
    doc = _parse_docx(buf)
    table = doc.tables[0]
    headers = [cell.text for cell in table.rows[0].cells]
    assert "Name" in headers
    assert "Email" in headers


def test_table_component_renders_rows() -> None:
    template = Template(
        id="t4",
        name="Test",
        output_format="docx",
        components=(
            TemplateComponent(
                id="c4",
                type="table",
                position={"x": 0, "y": 0, "width": 500, "height": "auto"},
                source_table="users",
                columns=({"header": "Name", "field": "name"},),
            ),
        ),
    )
    data = {"users": [{"name": "Alice"}, {"name": "Bob"}]}
    buf = render_word(template, data=data)
    doc = _parse_docx(buf)
    table = doc.tables[0]
    # Row 0 = headers, rows 1+ = data
    data_names = [table.rows[i].cells[0].text for i in range(1, len(table.rows))]
    assert "Alice" in data_names
    assert "Bob" in data_names
