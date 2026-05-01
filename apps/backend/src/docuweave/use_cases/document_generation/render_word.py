"""
Word (.docx) renderer.
Walks template components in sort order and writes python-docx content.
Returns an in-memory BytesIO buffer — no disk writes.
"""
import io
from typing import Any

from docx import Document
from docx.shared import Pt

from docuweave.domain.entities.template import Template, TemplateComponent
from docuweave.use_cases.transformations.execute_pipeline import execute_pipeline


def render_word(template: Template, data: dict[str, list[dict[str, Any]]]) -> io.BytesIO:
    doc = Document()

    for component in template.components:
        _render_component(doc, component, data)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _render_component(
    doc: Document,
    component: TemplateComponent,
    data: dict[str, list[dict[str, Any]]],
) -> None:
    if component.type == "text":
        doc.add_paragraph(component.content)

    elif component.type == "dynamic_field":
        rows = data.get(component.source_table, [])
        for row in rows:
            pipeline = list(component.transformations)
            if pipeline:
                value = execute_pipeline(row, pipeline, dataset=rows)
            else:
                value = row.get(component.source_field, "")
            doc.add_paragraph(str(value))

    elif component.type == "table":
        rows = data.get(component.source_table, [])
        columns = list(component.columns)
        if not columns:
            return

        table = doc.add_table(rows=1 + len(rows), cols=len(columns))
        table.style = "Table Grid"

        # Header row
        for col_idx, col_def in enumerate(columns):
            table.rows[0].cells[col_idx].text = col_def.get("header", "")

        # Data rows
        for row_idx, row in enumerate(rows, start=1):
            for col_idx, col_def in enumerate(columns):
                table.rows[row_idx].cells[col_idx].text = str(row.get(col_def.get("field", ""), ""))

    elif component.type == "repeat_block":
        rows = data.get(component.repeat_over, [])
        for row in rows:
            for sub in component.sub_components:
                content = sub.content
                # Simple token substitution: {{field_name}}
                for key, val in row.items():
                    content = content.replace(f"{{{{{key}}}}}", str(val))
                doc.add_paragraph(content)
